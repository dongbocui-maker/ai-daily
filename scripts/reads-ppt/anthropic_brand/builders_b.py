"""Builders B — two_column, insight_pair, action_list, synthesis, closing."""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from . import style as S
from .drawing import add_rect, add_textbox, add_paragraphs, add_hrule, fill_background
from .chrome import (
    apply_page_background, add_brand_mark, add_kicker, add_footer,
    add_standard_page_frame, add_title, add_subtitle,
)
from .builders_a import _group_size_table


# --- two_column value sizer ---
# Column content width ~5.65" — generous, but long phrases like
# "Slowed-down users" (17 chars) still need shrinking from 36pt.
_TWO_COL_VALUE_TABLE = [
    (8,  36),
    (14, 30),
    (20, 24),
    (28, 20),
    (999, 18),
]

# --- insight_pair header sizer ---
# Card content width ~5.85". Default 16pt. Headers like
# "Early-career squeeze" (20 chars) fit at 16pt; longer needs shrink.
_INSIGHT_HEADER_TABLE = [
    (22, 16),
    (32, 14),
    (44, 12),
    (999, 11),
]

# --- action_list header sizer ---
# Header column width ~3.5". Default 13pt. Long headers like
# "Re-design junior pathways" (25 chars) need shrink to avoid wrap-clip.
_ACTION_HEADER_TABLE = [
    (20, 13),
    (28, 12),
    (36, 11),
    (999, 10),
]

# --- synthesis header sizer ---
# 3-col content width ~3.85". Default 15pt.
_SYNTHESIS_HEADER_TABLE = [
    (18, 15),
    (26, 13),
    (36, 12),
    (999, 11),
]


# ---------- Two Column ----------
def build_two_column(slide, sd, page, total):
    add_standard_page_frame(slide, sd.get("kicker", ""), page, total)
    add_title(slide, sd.get("title", ""), y=Inches(1.50), size=Pt(30))
    add_subtitle(slide, sd.get("subtitle", ""), y=Inches(2.85))

    left = sd.get("left", {})
    right = sd.get("right", {})

    col_w = Inches(5.85)
    gap = Inches(0.23)
    x_left = S.M_LEFT
    x_right = x_left + col_w + gap
    top = Inches(3.90)
    col_h = Inches(2.55)

    # PARALLEL ELEMENT RULE: both columns share the same value font size
    v_size = _group_size_table([left, right], key="value", table=_TWO_COL_VALUE_TABLE)

    _draw_two_col(slide, x_left, top, col_w, col_h, left, v_size=v_size)
    _draw_two_col(slide, x_right, top, col_w, col_h, right, v_size=v_size)

    # callout strip
    callout = sd.get("callout")
    if callout:
        cy = Inches(6.55)
        add_hrule(slide, S.M_LEFT, cy, S.CONTENT_W, color=S.DARK,
                  height=Pt(1.2))
        add_textbox(
            slide, S.M_LEFT, cy + Inches(0.10), S.CONTENT_W, Inches(0.45),
            callout,
            font=S.FONT_BODY, font_size=Pt(12.5), italic=True, color=S.DARK,
            line_spacing=1.35, margin=Inches(0),
        )


def _draw_two_col(slide, x, y, w, h, col, *, v_size=36):
    color_name = col.get("color", "orange")
    accent = S.ACCENTS.get(color_name, S.ORANGE)

    # vertical rule
    add_rect(slide, x, y, Pt(2), h, fill=accent, line=accent)

    # tag · value (large)
    tag = col.get("tag", "")
    value = col.get("value", "")
    cur_y = y + Inches(0.00)
    add_textbox(
        slide, x + Inches(0.20), cur_y, w - Inches(0.20), Inches(0.30),
        tag.upper(),
        font=S.FONT_HEAD, font_size=Pt(11), bold=True, color=accent,
        char_spacing=S.CSP_KICKER, margin=Inches(0),
    )
    # value (large) — size injected from parallel-element sizer
    add_textbox(
        slide, x + Inches(0.20), cur_y + Inches(0.32),
        w - Inches(0.20), Inches(0.70),
        value,
        font=S.FONT_HEAD, font_size=Pt(v_size), bold=True, color=S.DARK,
        line_spacing=1.0, margin=Inches(0),
    )
    # header
    add_textbox(
        slide, x + Inches(0.20), cur_y + Inches(1.05),
        w - Inches(0.20), Inches(0.45),
        col.get("header", ""),
        font=S.FONT_HEAD, font_size=Pt(13), bold=True, color=S.DARK,
        line_spacing=1.20, margin=Inches(0),
    )

    # points
    points = col.get("points", [])
    if points:
        specs = []
        for p in points:
            specs.append({
                "runs": [
                    {"text": S.BULLET + "  ", "size": Pt(10),
                     "color": accent, "bold": True, "font": S.FONT_HEAD},
                    {"text": p, "size": Pt(10.5),
                     "color": S.DARK, "font": S.FONT_BODY},
                ],
                "line_spacing": 1.32,
                "space_after": Pt(4),
            })
        add_paragraphs(
            slide, x + Inches(0.20), cur_y + Inches(1.55),
            w - Inches(0.20), h - Inches(1.60),
            specs, margin=Inches(0),
        )


# ---------- Insight Pair ----------
def build_insight_pair(slide, sd, page, total):
    add_standard_page_frame(slide, sd.get("kicker", ""), page, total)
    add_title(slide, sd.get("title", ""), y=Inches(1.50), size=Pt(30))
    add_subtitle(slide, sd.get("subtitle", ""), y=Inches(2.85))

    cards = sd.get("cards", [])[:2]
    col_w = Inches(5.85)
    gap = Inches(0.23)
    x_left = S.M_LEFT
    x_right = x_left + col_w + gap
    top = Inches(3.95)
    col_h = Inches(2.95)

    # PARALLEL ELEMENT RULE: both cards share the same header font size
    h_size = _group_size_table(cards, key="header", table=_INSIGHT_HEADER_TABLE)

    for i, c in enumerate(cards):
        x = x_left if i == 0 else x_right
        _draw_insight(slide, x, top, col_w, col_h, c, h_size=h_size)


def _draw_insight(slide, x, y, w, h, c, *, h_size=16):
    color_name = c.get("color", "orange")
    accent = S.ACCENTS.get(color_name, S.ORANGE)
    pad = Inches(0.05)

    # top thick rule
    add_hrule(slide, x, y, Inches(2.0), color=accent, height=Pt(3))
    # tag
    add_textbox(
        slide, x, y + Inches(0.12), w, Inches(0.30),
        c.get("tag", "").upper(),
        font=S.FONT_HEAD, font_size=Pt(10), bold=True, color=accent,
        char_spacing=S.CSP_KICKER, margin=Inches(0),
    )
    # header — size injected from parallel-element sizer
    add_textbox(
        slide, x, y + Inches(0.45), w, Inches(0.55),
        c.get("header", ""),
        font=S.FONT_HEAD, font_size=Pt(h_size), bold=True, color=S.DARK,
        line_spacing=1.15, margin=Inches(0),
    )
    # body
    add_textbox(
        slide, x, y + Inches(1.05), w, Inches(0.85),
        c.get("body", ""),
        font=S.FONT_BODY, font_size=Pt(11), color=S.DARK,
        line_spacing=1.40, margin=Inches(0),
    )

    # evidence — italic muted
    ev = c.get("evidence", "")
    ev_label = c.get("evidence_label", "Evidence")
    if ev:
        add_textbox(
            slide, x, y + Inches(1.92), w, Inches(0.25),
            ev_label.upper(),
            font=S.FONT_HEAD, font_size=Pt(9), bold=True, color=accent,
            char_spacing=S.CSP_LABEL, margin=Inches(0),
        )
        add_textbox(
            slide, x, y + Inches(2.18), w, Inches(0.45),
            ev,
            font=S.FONT_BODY, font_size=Pt(10), italic=True, color=S.MID_GRAY,
            line_spacing=1.30, margin=Inches(0),
        )

    # implication
    impl = c.get("implication", "")
    impl_label = c.get("implication_label", "Implication")
    if impl:
        add_paragraphs(
            slide, x, y + h - Inches(0.45), w, Inches(0.45),
            [{"runs": [
                {"text": impl_label + S.SEP, "size": Pt(10),
                 "bold": True, "color": accent, "font": S.FONT_HEAD},
                {"text": impl, "size": Pt(11),
                 "italic": True, "color": S.DARK, "font": S.FONT_BODY},
            ], "line_spacing": 1.32}],
            margin=Inches(0),
        )


# ---------- Action List ----------
# --- action_list what/value body sizer (shrinks long bodies to fit row_h) ---
_ACTION_BODY_TABLE = [
    (90,  11),   # short — default 11pt
    (120, 10),
    (160, 9),
    (999, 8),
]

_ACTION_VALUE_TABLE = [
    (60,  10),
    (85,  9),
    (120, 8),
    (999, 7.5),
]


def build_action_list(slide, sd, page, total):
    add_standard_page_frame(slide, sd.get("kicker", ""), page, total)
    add_title(slide, sd.get("title", ""), y=Inches(1.50), size=Pt(30))
    add_subtitle(slide, sd.get("subtitle", ""), y=Inches(2.85))

    actions = sd.get("actions", [])
    n = len(actions)
    if n == 0:
        return

    # vertical stack — 5 rows; lift top + extend bottom for more breathing room
    top = Inches(3.65)
    avail = Inches(7.00) - top   # 3.35" total → 0.67"/row for n=5
    row_h = avail / n

    # PARALLEL ELEMENT RULE: all action headers share the same font size
    head_size = _group_size_table(actions, key="header", table=_ACTION_HEADER_TABLE)
    # PARALLEL ELEMENT RULE: all action `what` and `value` bodies share one size
    what_size = _group_size_table(actions, key="what",  table=_ACTION_BODY_TABLE)
    val_size  = _group_size_table(actions, key="value", table=_ACTION_VALUE_TABLE)

    for i, act in enumerate(actions):
        y = top + row_h * i
        color_name = act.get("color", "orange")
        accent = S.ACCENTS.get(color_name, S.ORANGE)
        _draw_action_row(slide, y, row_h, act, accent,
                         head_size=head_size,
                         what_size=what_size,
                         val_size=val_size)


def _draw_action_row(slide, y, h, act, accent, *, head_size=13,
                     what_size=11, val_size=10):
    # number column
    num_x = S.M_LEFT
    num_w = Inches(0.85)
    add_textbox(
        slide, num_x, y, num_w, h,
        act.get("num", ""),
        font=S.FONT_HEAD, font_size=Pt(26), bold=True, color=accent,
        anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0),
    )
    # accent vertical line after num
    add_rect(slide, num_x + num_w, y + Inches(0.10),
             Pt(1.5), h - Inches(0.20),
             fill=accent, line=accent)
    # action header column — size injected from parallel-element sizer
    head_x = num_x + num_w + Inches(0.30)
    head_w = Inches(3.5)
    add_textbox(
        slide, head_x, y, head_w, h,
        act.get("header", ""),
        font=S.FONT_HEAD, font_size=Pt(head_size), bold=True, color=S.DARK,
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.20, margin=Inches(0),
    )
    # what to do column — size injected from parallel-element sizer
    what_x = head_x + head_w + Inches(0.20)
    what_w = Inches(4.3)
    add_textbox(
        slide, what_x, y, what_w, h,
        act.get("what", ""),
        font=S.FONT_BODY, font_size=Pt(what_size), color=S.DARK,
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.28, margin=Inches(0),
    )
    # value column — size injected from parallel-element sizer
    val_x = what_x + what_w + Inches(0.20)
    val_w = Inches(2.50)
    add_textbox(
        slide, val_x, y, val_w, h,
        act.get("value", ""),
        font=S.FONT_BODY, font_size=Pt(val_size), italic=True, color=accent,
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.28, margin=Inches(0),
    )
    # bottom rule
    add_hrule(slide, S.M_LEFT, y + h - Pt(0.6), S.CONTENT_W,
              color=S.LIGHT_GRAY, height=Pt(0.5))


# ---------- Synthesis ----------
def build_synthesis(slide, sd, page, total):
    add_standard_page_frame(slide, sd.get("kicker", ""), page, total)
    add_title(slide, sd.get("title", ""), y=Inches(1.50), size=Pt(30))
    add_subtitle(slide, sd.get("subtitle", ""), y=Inches(2.85))

    pillars = sd.get("pillars", [])[:3]
    n = len(pillars)
    if n == 0:
        return

    total_w = S.CONTENT_W
    gap = Inches(0.30)
    col_w = (total_w - gap * (n - 1)) / n
    top = Inches(3.85)
    col_h = Inches(2.65)

    # PARALLEL ELEMENT RULE: all 3 pillar headers share the same font size
    p_head_size = _group_size_table(pillars, key="header", table=_SYNTHESIS_HEADER_TABLE)

    for i, p in enumerate(pillars):
        x = S.M_LEFT + (col_w + gap) * i
        color_name = p.get("color", "orange")
        accent = S.ACCENTS.get(color_name, S.ORANGE)
        # top accent rule
        add_hrule(slide, x, top, Inches(2.5), color=accent, height=Pt(3))
        # num
        add_textbox(
            slide, x, top + Inches(0.20), col_w, Inches(0.65),
            p.get("num", ""),
            font=S.FONT_HEAD, font_size=Pt(36), bold=True, color=accent,
            margin=Inches(0),
        )
        # header — size injected from parallel-element sizer
        add_textbox(
            slide, x, top + Inches(0.95), col_w, Inches(0.5),
            p.get("header", ""),
            font=S.FONT_HEAD, font_size=Pt(p_head_size), bold=True, color=S.DARK,
            line_spacing=1.20, margin=Inches(0),
        )
        # body
        add_textbox(
            slide, x, top + Inches(1.55), col_w, col_h - Inches(1.60),
            p.get("body", ""),
            font=S.FONT_BODY, font_size=Pt(11), color=S.DARK,
            line_spacing=1.42, margin=Inches(0),
        )

    closing = sd.get("closing_line")
    if closing:
        cy = Inches(6.65)
        add_hrule(slide, S.M_LEFT, cy, S.CONTENT_W, color=S.DARK,
                  height=Pt(1.2))
        add_textbox(
            slide, S.M_LEFT, cy + Inches(0.10), S.CONTENT_W, Inches(0.40),
            closing,
            font=S.FONT_BODY, font_size=Pt(13), italic=True, color=S.DARK,
            line_spacing=1.35, margin=Inches(0),
        )


# ---------- Closing ----------
def build_closing(slide, sd, page, total):
    apply_page_background(slide, color=S.DARK)
    add_brand_mark(slide, on_dark=True)
    add_footer(slide, page, total, on_dark=True)

    add_kicker(slide, sd.get("kicker", ""), y=Inches(1.50),
               color=S.ORANGE)
    # title
    add_textbox(
        slide, S.M_LEFT, Inches(2.10), Inches(12.0), Inches(1.30),
        sd.get("title", ""),
        font=S.FONT_HEAD, font_size=Pt(72), bold=True, color=S.LIGHT,
        line_spacing=1.0, margin=Inches(0),
    )
    sub = sd.get("subtitle", "")
    if sub:
        add_textbox(
            slide, S.M_LEFT, Inches(3.55), Inches(12.0), Inches(0.5),
            sub,
            font=S.FONT_BODY, font_size=Pt(14), italic=True, color=S.MID_GRAY,
            margin=Inches(0),
        )

    # primary link block
    link = sd.get("primary_link", {})
    if link:
        ly = Inches(4.85)
        add_hrule(slide, S.M_LEFT, ly, Inches(2.5),
                  color=S.ORANGE, height=Pt(2))
        add_textbox(
            slide, S.M_LEFT, ly + Inches(0.15), Inches(12.0), Inches(0.35),
            "READ THE FULL REPORT",
            font=S.FONT_HEAD, font_size=Pt(10), bold=True, color=S.ORANGE,
            char_spacing=S.CSP_KICKER, margin=Inches(0),
        )
        add_textbox(
            slide, S.M_LEFT, ly + Inches(0.55), Inches(12.0), Inches(0.50),
            link.get("label", ""),
            font=S.FONT_HEAD, font_size=Pt(20), bold=True, color=S.LIGHT,
            margin=Inches(0),
        )
        add_textbox(
            slide, S.M_LEFT, ly + Inches(1.05), Inches(12.0), Inches(0.40),
            link.get("url", ""),
            font=S.FONT_BODY, font_size=Pt(11), italic=True, color=S.MID_GRAY,
            margin=Inches(0),
        )

    # authors line at bottom
    al = sd.get("authors_line", "")
    if al:
        add_textbox(
            slide, S.M_LEFT, Inches(6.50), S.CONTENT_W, Inches(0.45),
            al,
            font=S.FONT_BODY, font_size=Pt(9.5), italic=True, color=S.MID_GRAY,
            line_spacing=1.35, margin=Inches(0),
        )
