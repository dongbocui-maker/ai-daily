"""Builders B — stat_grid, three_card, two_column, quote_pair."""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from . import style as S
from .drawing import add_rect, add_textbox, add_paragraphs, add_card
from .chrome import add_header, add_footer


def _fit_size(text, *, base=36, max_chars=7, min_size=14, step=4):
    if not text:
        return base
    L = len(text)
    if L <= max_chars:
        return base
    if L <= 10:
        return max(min_size, base - step)
    if L <= 14:
        return max(min_size, base - step * 3)
    if L <= 18:
        return max(min_size, base - step * 5)
    return min_size


# ---------- Stat Grid ----------
def build_stat_grid(slide, sd, page, total):
    add_header(slide, sd.get("eyebrow", ""), sd.get("title", ""),
               sd.get("subtitle"))
    add_footer(slide, page, total)

    stats = sd.get("stats", [])
    n = len(stats)
    if n == 0:
        return

    total_w = Inches(12.333)
    gap = Inches(0.15)
    card_w = (total_w - gap * (n - 1)) / n
    card_h = Inches(3.6)
    top = Inches(2.20)

    for i, st in enumerate(stats):
        x = Inches(0.5) + (card_w + gap) * i
        add_card(slide, x, top, card_w, card_h, accent_top=True)
        # value (auto-shrink long values)
        v = st.get("value", "")
        v_size = _fit_size(v, base=36, max_chars=7, min_size=20)
        add_textbox(
            slide, x, top + Inches(0.30), card_w, Inches(0.95),
            v,
            font_size=Pt(v_size), bold=True, color=S.PURPLE_DEEP,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0.05),
        )
        rule_w = Inches(0.6)
        add_rect(slide, x + (card_w - rule_w) / 2, top + Inches(1.42),
                 rule_w, Inches(0.025), fill=S.PURPLE, line=S.PURPLE)
        # label
        add_textbox(
            slide, x + Inches(0.15), top + Inches(1.58),
            card_w - Inches(0.3), Inches(0.6),
            st.get("label", ""),
            font_size=Pt(11), bold=True, color=S.INK,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP,
            line_spacing=1.20, margin=Inches(0.02),
        )
        # context paragraph
        ctx = st.get("context", "")
        if ctx:
            add_textbox(
                slide, x + Inches(0.15), top + Inches(2.20),
                card_w - Inches(0.3), Inches(1.0),
                ctx,
                font_size=Pt(8.5), italic=True, color=S.INK_SOFT,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP,
                line_spacing=1.25, margin=Inches(0.02),
            )
        # source line at bottom
        src = st.get("source", "")
        if src:
            add_textbox(
                slide, x + Inches(0.15), top + card_h - Inches(0.30),
                card_w - Inches(0.3), Inches(0.22),
                src,
                font_size=Pt(7.5), italic=True, color=S.MUTED,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0),
            )

    # narrative strip at bottom
    narrative = sd.get("narrative")
    if narrative:
        strip_y = Inches(6.30)
        add_rect(slide, Inches(0.5), strip_y, Inches(12.333), Inches(0.60),
                 fill=S.INK, line=S.INK)
        add_rect(slide, Inches(0.5), strip_y, Inches(0.10), Inches(0.60),
                 fill=S.PURPLE, line=S.PURPLE)
        add_textbox(
            slide, Inches(0.75), strip_y, Inches(11.9), Inches(0.60),
            narrative,
            font_size=Pt(10.5), italic=True, color=S.WHITE,
            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.25, margin=Inches(0.05),
        )


# ---------- Three Card ----------
def build_three_card(slide, sd, page, total):
    add_header(slide, sd.get("eyebrow", ""), sd.get("title", ""),
               sd.get("subtitle"))
    add_footer(slide, page, total)

    cards = sd.get("cards", [])
    n = len(cards)
    if n == 0:
        return

    total_w = Inches(12.333)
    gap = Inches(0.20)
    card_w = (total_w - gap * (n - 1)) / n
    card_h = Inches(4.80)
    top = Inches(2.15)

    for i, c in enumerate(cards):
        x = Inches(0.5) + (card_w + gap) * i
        _draw_three_card(slide, x, top, card_w, card_h, c)


def _draw_three_card(slide, x, y, w, h, c):
    add_card(slide, x, y, w, h, accent_top=True)
    pad = Inches(0.25)
    inner_w = w - pad * 2
    cur_y = y + Inches(0.22)

    # tag (PATTERN 01 etc)
    tag = c.get("tag")
    if tag:
        add_textbox(
            slide, x + pad, cur_y, inner_w, Inches(0.25),
            tag,
            font_size=Pt(10), bold=True, color=S.PURPLE,
            char_spacing=S.CSP_PILL, margin=Inches(0),
        )
        cur_y = cur_y + Inches(0.30)

    # header (card title)
    add_textbox(
        slide, x + pad, cur_y, inner_w, Inches(0.85),
        c.get("header", ""),
        font_size=Pt(13.5), bold=True, color=S.INK,
        line_spacing=1.20, margin=Inches(0),
    )
    cur_y = cur_y + Inches(0.90)

    # purple rule
    add_rect(slide, x + pad, cur_y, Inches(0.5), Inches(0.025),
             fill=S.PURPLE, line=S.PURPLE)
    cur_y = cur_y + Inches(0.15)

    # body
    body = c.get("body", "")
    add_textbox(
        slide, x + pad, cur_y, inner_w, Inches(1.30),
        body,
        font_size=Pt(10), color=S.INK_SOFT,
        line_spacing=1.30, margin=Inches(0),
    )
    cur_y = cur_y + Inches(1.40)

    # evidence block (with label)
    evidence = c.get("evidence")
    ev_label = c.get("evidence_label", "Evidence")
    if evidence:
        # small tinted box
        ev_y = y + h - Inches(1.40)
        ev_h = Inches(1.10)
        add_rect(slide, x + pad, ev_y, inner_w, ev_h,
                 fill=S.PURPLE_TINT, line=S.PURPLE_EDGE)
        add_paragraphs(
            slide, x + pad + Inches(0.10), ev_y + Inches(0.10),
            inner_w - Inches(0.2), ev_h - Inches(0.2),
            [
                {"runs": [
                    {"text": ev_label.upper() + "  ", "size": Pt(8),
                     "bold": True, "color": S.PURPLE,
                     "char_spacing": S.CSP_PILL},
                ], "space_after": Pt(2)},
                {"runs": [
                    {"text": evidence, "size": Pt(9),
                     "italic": True, "color": S.PURPLE_DEEP},
                ], "line_spacing": 1.30},
            ],
            margin=Inches(0),
        )

    # source at very bottom
    src = c.get("source")
    if src:
        add_textbox(
            slide, x + pad, y + h - Inches(0.22), inner_w, Inches(0.18),
            "— " + src,
            font_size=Pt(7.5), italic=True, color=S.MUTED,
            anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0),
        )


# ---------- Two Column ----------
def build_two_column(slide, sd, page, total):
    add_header(slide, sd.get("eyebrow", ""), sd.get("title", ""),
               sd.get("subtitle"))
    add_footer(slide, page, total)

    left = sd.get("left", {})
    right = sd.get("right", {})

    col_w = Inches(6.0)
    gap = Inches(0.333)
    col_h = Inches(3.75)
    x_left = Inches(0.5)
    x_right = x_left + col_w + gap
    top = Inches(2.15)

    _draw_two_col_card(slide, x_left, top, col_w, col_h, left)
    _draw_two_col_card(slide, x_right, top, col_w, col_h, right)

    # callout strip
    callout = sd.get("callout")
    callout_label = sd.get("callout_label", "Why this matters")
    if callout:
        co_y = top + col_h + Inches(0.20)
        co_h = Inches(0.85)
        add_rect(slide, Inches(0.5), co_y, Inches(12.333), co_h,
                 fill=S.INK, line=S.INK)
        add_rect(slide, Inches(0.5), co_y, Inches(0.12), co_h,
                 fill=S.PURPLE, line=S.PURPLE)
        add_paragraphs(
            slide, Inches(0.75), co_y, Inches(11.9), co_h,
            [
                {"runs": [
                    {"text": callout_label.upper() + "  ·  ",
                     "size": Pt(10), "bold": True, "color": S.PURPLE,
                     "char_spacing": S.CSP_PILL},
                    {"text": callout, "size": Pt(11),
                     "italic": True, "color": S.WHITE},
                ], "line_spacing": 1.30},
            ],
            margin=Inches(0.10),
            anchor=MSO_ANCHOR.MIDDLE,
        )

    # source note (small)
    sn = sd.get("source_note")
    if sn:
        add_textbox(
            slide, Inches(0.5), Inches(7.0), Inches(12.333), Inches(0.20),
            sn,
            font_size=Pt(7.5), italic=True, color=S.MUTED,
            anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0),
        )


def _draw_two_col_card(slide, x, y, w, h, col):
    add_card(slide, x, y, w, h, accent_top=True)
    pad = Inches(0.28)
    cur_y = y + Inches(0.22)

    label = col.get("label", "")
    if label:
        add_textbox(
            slide, x + pad, cur_y, w - pad * 2, Inches(0.28),
            label,
            font_size=Pt(10), bold=True, color=S.PURPLE,
            char_spacing=S.CSP_PILL, margin=Inches(0),
        )
        cur_y = cur_y + Inches(0.32)

    header = col.get("header", "")
    add_textbox(
        slide, x + pad, cur_y, w - pad * 2, Inches(0.55),
        header,
        font_size=Pt(14), bold=True, color=S.INK,
        line_spacing=1.20, margin=Inches(0),
    )
    cur_y = cur_y + Inches(0.60)

    add_rect(slide, x + pad, cur_y, Inches(0.5), Inches(0.025),
             fill=S.PURPLE, line=S.PURPLE)
    cur_y = cur_y + Inches(0.15)

    points = col.get("points", [])
    if points:
        specs = []
        for p in points:
            specs.append({
                "runs": [
                    {"text": S.BULLET_CHAR + "  ", "size": Pt(9.5),
                     "color": S.PURPLE, "bold": True},
                    {"text": p, "size": Pt(10), "color": S.INK_SOFT},
                ],
                "line_spacing": 1.30,
                "space_after": Pt(5),
            })
        add_paragraphs(
            slide, x + pad, cur_y, w - pad * 2, y + h - cur_y - Inches(0.15),
            specs, margin=Inches(0),
        )


# ---------- Quote Pair (primary + secondary quote + implication) ----------
def build_quote_pair(slide, sd, page, total):
    add_header(slide, sd.get("eyebrow", ""), sd.get("title", ""),
               sd.get("subtitle"))
    add_footer(slide, page, total)

    primary = sd.get("primary_quote", "")
    primary_attrib = sd.get("primary_attrib", "")
    secondary = sd.get("secondary_quote", "")
    secondary_attrib = sd.get("secondary_attrib", "")
    implication = sd.get("implication", "")
    impl_label = sd.get("implication_label", "Why this matters")

    # Layout: 2 quote cards side by side + bottom implication strip
    col_w = Inches(6.0)
    gap = Inches(0.333)
    col_h = Inches(3.75)
    x_left = Inches(0.5)
    x_right = x_left + col_w + gap
    top = Inches(2.15)

    _draw_quote(slide, x_left, top, col_w, col_h,
                primary, primary_attrib, primary=True)
    _draw_quote(slide, x_right, top, col_w, col_h,
                secondary, secondary_attrib, primary=False)

    # implication strip
    if implication:
        im_y = top + col_h + Inches(0.20)
        im_h = Inches(0.90)
        add_rect(slide, Inches(0.5), im_y, Inches(12.333), im_h,
                 fill=S.INK, line=S.INK)
        add_rect(slide, Inches(0.5), im_y, Inches(0.12), im_h,
                 fill=S.PURPLE, line=S.PURPLE)
        add_paragraphs(
            slide, Inches(0.75), im_y, Inches(11.9), im_h,
            [
                {"runs": [
                    {"text": impl_label.upper() + "  ·  ",
                     "size": Pt(10), "bold": True, "color": S.PURPLE,
                     "char_spacing": S.CSP_PILL},
                    {"text": implication, "size": Pt(11),
                     "italic": True, "color": S.WHITE},
                ], "line_spacing": 1.30},
            ],
            margin=Inches(0.10),
            anchor=MSO_ANCHOR.MIDDLE,
        )


def _draw_quote(slide, x, y, w, h, text, attrib, primary=True):
    accent = S.PURPLE if primary else S.TEAL
    add_card(slide, x, y, w, h, accent_top=True, accent_color=accent)
    # giant left quote glyph
    add_textbox(
        slide, x + Inches(0.20), y + Inches(0.12),
        Inches(1.0), Inches(0.9),
        "\u201C",
        font_size=Pt(56), bold=True, color=accent,
        anchor=MSO_ANCHOR.TOP, margin=Inches(0),
    )
    # body
    body_y = y + Inches(0.95)
    body_h = h - Inches(1.55)
    add_textbox(
        slide, x + Inches(0.35), body_y, w - Inches(0.7), body_h,
        text,
        font_size=Pt(13), italic=True, color=S.INK_SOFT,
        line_spacing=1.40, margin=Inches(0),
    )
    # attribution
    if attrib:
        add_textbox(
            slide, x + Inches(0.35), y + h - Inches(0.55),
            w - Inches(0.7), Inches(0.40),
            "— " + attrib,
            font_size=Pt(9), color=S.MUTED, margin=Inches(0),
        )
