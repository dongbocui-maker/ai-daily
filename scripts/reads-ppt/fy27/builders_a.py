"""Slide builders: cover, context, agenda, section_header, stat_grid."""

from pptx.util import Inches, Pt

from .style import C, SLIDE_H
from .drawing import (
    add_blank_slide, set_bg, add_rect, add_textbox, add_paragraphs,
    add_header, add_footer,
)


def build_cover(prs, data, total, page_no):
    slide = add_blank_slide(prs)
    set_bg(slide)
    # left purple strip
    add_rect(slide, Inches(0), Inches(0), Inches(0.4), SLIDE_H, fill=C["purple"])
    # accenture wordmark
    add_textbox(slide, Inches(0.8), Inches(0.6), Inches(2.0), Inches(0.4),
                "accenture", font_size=14, bold=True, color=C["ink"], valign="middle")
    add_rect(slide, Inches(2.45), Inches(0.74), Inches(0.16), Inches(0.16), fill=C["purple"])

    add_textbox(slide, Inches(0.8), Inches(2.6), Inches(11.5), Inches(0.4),
                data.get("eyebrow", "AI DAILY READS").upper(),
                font_size=12, bold=True, color=C["purple"], char_spacing=5)
    add_textbox(slide, Inches(0.8), Inches(3.05), Inches(11.5), Inches(1.5),
                data.get("title", ""), font_size=44, bold=True, color=C["ink"])
    add_textbox(slide, Inches(0.8), Inches(4.5), Inches(11.5), Inches(1.0),
                data.get("subtitle", ""), font_size=16, italic=True, color=C["ink_soft"])

    if data.get("source_line"):
        add_textbox(slide, Inches(0.8), Inches(5.6), Inches(11.5), Inches(0.3),
                    data["source_line"], font_size=10, color=C["muted"])

    highlights = data.get("highlights", [])
    if highlights:
        n = len(highlights)
        total_w = 11.5
        gap = 0.2
        tile_w = (total_w - gap * (n - 1)) / n
        top = 6.15
        h = 0.85
        for i, item in enumerate(highlights):
            x = 0.8 + i * (tile_w + gap)
            add_rect(slide, Inches(x), Inches(top), Inches(tile_w), Inches(h),
                     fill=C["purple_tint"], line=C["purple_edge"], line_w=Pt(0.5))
            add_rect(slide, Inches(x), Inches(top), Inches(0.08), Inches(h), fill=C["purple"])
            add_textbox(slide, Inches(x + 0.18), Inches(top + 0.08),
                        Inches(tile_w - 0.25), Inches(0.25),
                        item.get("label", "").upper(), font_size=8.5, bold=True,
                        color=C["purple_deep"], char_spacing=2)
            add_textbox(slide, Inches(x + 0.18), Inches(top + 0.32),
                        Inches(tile_w - 0.25), Inches(0.5),
                        item.get("value", ""), font_size=16, bold=True, color=C["ink"])


def build_context(prs, data, total, page_no):
    slide = add_blank_slide(prs)
    set_bg(slide)
    add_header(slide, eyebrow=data.get("eyebrow"), title=data.get("title"),
               subtitle=data.get("subtitle"))

    if data.get("key_message"):
        add_rect(slide, Inches(0.5), Inches(2.05), Inches(12.3), Inches(0.45),
                 fill=C["purple_tint"], line=C["purple_edge"], line_w=Pt(0.5))
        add_rect(slide, Inches(0.5), Inches(2.05), Inches(0.08), Inches(0.45), fill=C["purple"])
        add_textbox(slide, Inches(0.72), Inches(2.10), Inches(12), Inches(0.35),
                    data["key_message"], font_size=11.5, italic=True,
                    color=C["ink"], valign="middle")

    col_top = 2.75
    col_h = 3.4
    col_w = 5.9
    gap = 0.5
    col_x = [0.5, 0.5 + col_w + gap]
    blocks = [data.get("from_block", {}), data.get("to_block", {})]

    for i, blk in enumerate(blocks):
        x = col_x[i]
        add_rect(slide, Inches(x), Inches(col_top), Inches(col_w), Inches(col_h),
                 fill=C["white"], line=C["border"], line_w=Pt(0.5))
        add_rect(slide, Inches(x), Inches(col_top), Inches(col_w), Inches(0.05), fill=C["purple"])

        label = blk.get("label", "FROM" if i == 0 else "TO")
        label_color = C["muted"] if i == 0 else C["purple"]
        add_textbox(slide, Inches(x + 0.25), Inches(col_top + 0.18),
                    Inches(2), Inches(0.3),
                    label, font_size=10, bold=True, color=label_color, char_spacing=3)
        add_textbox(slide, Inches(x + 0.25), Inches(col_top + 0.50),
                    Inches(col_w - 0.5), Inches(0.45),
                    blk.get("header", ""), font_size=15, bold=True, color=C["ink"])
        add_paragraphs(slide, Inches(x + 0.25), Inches(col_top + 1.05),
                       Inches(col_w - 0.5), Inches(col_h - 1.2),
                       blk.get("points", []),
                       font_size=10.5, color=C["ink_soft"], bullet=True,
                       line_spacing=1.35)

    stat_band = data.get("stat_band", [])
    if stat_band:
        band_top = 6.30
        n = len(stat_band)
        gap = 0.25
        tile_w = (12.3 - gap * (n - 1)) / n
        for i, item in enumerate(stat_band):
            x = 0.5 + i * (tile_w + gap)
            add_rect(slide, Inches(x), Inches(band_top), Inches(tile_w), Inches(0.55),
                     fill=C["bg_soft"], line=C["border"], line_w=Pt(0.5))
            add_textbox(slide, Inches(x), Inches(band_top + 0.05),
                        Inches(tile_w), Inches(0.27),
                        item.get("value", ""), font_size=14, bold=True,
                        color=C["purple_deep"], align="center")
            add_textbox(slide, Inches(x), Inches(band_top + 0.32),
                        Inches(tile_w), Inches(0.2),
                        item.get("label", ""), font_size=8.5, color=C["muted"], align="center")

    if data.get("source_note"):
        add_textbox(slide, Inches(0.5), Inches(6.92), Inches(12.3), Inches(0.2),
                    data["source_note"], font_size=8, italic=True, color=C["muted"])

    add_footer(slide, page_no, total)


def build_agenda(prs, data, total, page_no):
    slide = add_blank_slide(prs)
    set_bg(slide)
    add_header(slide, eyebrow=data.get("eyebrow"), title=data.get("title"),
               subtitle=data.get("subtitle"))

    items = data.get("items", [])
    n = max(len(items), 1)
    row_top = 2.4
    row_gap = 0.2
    row_h = (7.0 - row_top - 0.2 - (n - 1) * row_gap) / n

    for i, item in enumerate(items):
        y = row_top + i * (row_h + row_gap)
        add_rect(slide, Inches(0.5), Inches(y), Inches(12.3), Inches(row_h),
                 fill=C["white"], line=C["border"], line_w=Pt(0.5))
        add_rect(slide, Inches(0.5), Inches(y), Inches(0.08), Inches(row_h), fill=C["purple"])
        num = item.get("num", str(i + 1).zfill(2))
        add_textbox(slide, Inches(0.75), Inches(y + 0.1), Inches(1.1), Inches(row_h - 0.2),
                    num, font_size=24, bold=True, color=C["purple"], valign="middle")
        add_textbox(slide, Inches(2.0), Inches(y + 0.10), Inches(10.5), Inches(0.4),
                    item.get("label", ""), font_size=15, bold=True, color=C["ink"])
        add_textbox(slide, Inches(2.0), Inches(y + 0.46), Inches(10.5), Inches(0.4),
                    item.get("sub", ""), font_size=10.5, italic=True, color=C["ink_soft"])

    add_footer(slide, page_no, total)


def build_section_header(prs, data, total, page_no):
    slide = add_blank_slide(prs)
    set_bg(slide)
    add_rect(slide, Inches(0), Inches(0), Inches(4.5), SLIDE_H, fill=C["purple"])

    add_textbox(slide, Inches(0.5), Inches(1.4), Inches(3.5), Inches(0.5),
                data.get("eyebrow", "").upper(), font_size=14, bold=True,
                color=C["white"], char_spacing=5)
    add_textbox(slide, Inches(0.5), Inches(2.0), Inches(3.7), Inches(2.5),
                data.get("title", ""), font_size=40, bold=True, color=C["white"])

    add_textbox(slide, Inches(5.0), Inches(2.0), Inches(7.8), Inches(0.6),
                data.get("subtitle", ""), font_size=18, italic=True, color=C["ink_soft"])

    preview = data.get("preview", [])
    if preview:
        add_textbox(slide, Inches(5.0), Inches(3.2), Inches(7.8), Inches(0.3),
                    "WHAT'S IN THIS ACT", font_size=10, bold=True,
                    color=C["purple"], char_spacing=4)
        add_paragraphs(slide, Inches(5.0), Inches(3.6), Inches(7.8), Inches(3),
                       preview, font_size=14, color=C["ink"], bullet=True,
                       line_spacing=1.5)

    add_footer(slide, page_no, total)


def build_stat_grid(prs, data, total, page_no):
    slide = add_blank_slide(prs)
    set_bg(slide)
    add_header(slide, eyebrow=data.get("eyebrow"), title=data.get("title"),
               subtitle=data.get("subtitle"))

    stats = data.get("stats", [])
    n = len(stats)
    cols = 2
    rows = (n + cols - 1) // cols
    grid_top = 2.20
    grid_left = 0.5
    grid_w = 12.3
    grid_h = 3.8
    gap = 0.25
    cell_w = (grid_w - gap * (cols - 1)) / cols
    cell_h = (grid_h - gap * (rows - 1)) / rows

    for i, stat in enumerate(stats):
        r = i // cols
        c = i % cols
        x = grid_left + c * (cell_w + gap)
        y = grid_top + r * (cell_h + gap)
        add_rect(slide, Inches(x), Inches(y), Inches(cell_w), Inches(cell_h),
                 fill=C["white"], line=C["border"], line_w=Pt(0.5))
        add_rect(slide, Inches(x), Inches(y), Inches(cell_w), Inches(0.05), fill=C["purple"])

        add_textbox(slide, Inches(x + 0.3), Inches(y + 0.18),
                    Inches(cell_w - 0.6), Inches(0.7),
                    stat.get("value", ""), font_size=32, bold=True, color=C["purple_deep"])
        add_textbox(slide, Inches(x + 0.3), Inches(y + 0.92),
                    Inches(cell_w - 0.6), Inches(0.4),
                    stat.get("label", ""), font_size=11.5, bold=True, color=C["ink"])
        add_textbox(slide, Inches(x + 0.3), Inches(y + 1.30),
                    Inches(cell_w - 0.6), Inches(0.55),
                    stat.get("context", ""), font_size=9.5, italic=True, color=C["ink_soft"])
        if stat.get("source"):
            add_textbox(slide, Inches(x + 0.3), Inches(y + cell_h - 0.30),
                        Inches(cell_w - 0.6), Inches(0.25),
                        f"Source: {stat['source']}", font_size=8,
                        italic=True, color=C["muted"])

    if data.get("narrative"):
        nar_top = grid_top + grid_h + 0.15
        add_rect(slide, Inches(0.5), Inches(nar_top), Inches(12.3), Inches(0.7),
                 fill=C["purple_tint"], line=C["purple_edge"], line_w=Pt(0.5))
        add_rect(slide, Inches(0.5), Inches(nar_top), Inches(0.08), Inches(0.7),
                 fill=C["purple"])
        add_textbox(slide, Inches(0.72), Inches(nar_top + 0.10),
                    Inches(12), Inches(0.55),
                    data["narrative"], font_size=10.5, italic=True,
                    color=C["ink_soft"], valign="middle")

    add_footer(slide, page_no, total)
