"""Builders C — insight_pair, action_table, summary, references, closing."""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from . import style as S
from .drawing import add_rect, add_textbox, add_paragraphs, add_card, add_pill
from .chrome import add_header, add_footer


# ---------- Insight Pair (cards: tag/header/body/evidence/implication) ----------
def build_insight_pair(slide, sd, page, total):
    add_header(slide, sd.get("eyebrow", ""), sd.get("title", ""),
               sd.get("subtitle"))
    add_footer(slide, page, total)

    cards = sd.get("cards", sd.get("insights", []))[:2]
    if not cards:
        return

    col_w = Inches(6.0)
    gap = Inches(0.333)
    col_h = Inches(4.65)
    x_left = Inches(0.5)
    x_right = x_left + col_w + gap
    top = Inches(2.15)

    for i, c in enumerate(cards):
        x = x_left if i == 0 else x_right
        _draw_insight_card(slide, x, top, col_w, col_h, c)

    sn = sd.get("source_note")
    if sn:
        add_textbox(
            slide, Inches(0.5), Inches(6.95), Inches(12.333), Inches(0.20),
            sn,
            font_size=Pt(8), italic=True, color=S.MUTED,
            anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0),
        )


def _draw_insight_card(slide, x, y, w, h, c):
    add_card(slide, x, y, w, h, accent_top=True)
    pad = Inches(0.28)
    cur_y = y + Inches(0.22)
    inner_w = w - pad * 2

    # pill tag
    tag = c.get("tag", "INSIGHT")
    add_pill(slide, x + pad, cur_y, Inches(0.32),
             tag, pad_x=Inches(0.18))
    cur_y = cur_y + Inches(0.50)

    # header
    add_textbox(
        slide, x + pad, cur_y, inner_w, Inches(0.7),
        c.get("header", ""),
        font_size=Pt(15), bold=True, color=S.INK,
        line_spacing=1.18, margin=Inches(0),
    )
    cur_y = cur_y + Inches(0.78)

    # body
    body = c.get("body", "")
    body_h = Inches(1.10)
    add_textbox(
        slide, x + pad, cur_y, inner_w, body_h,
        body,
        font_size=Pt(10.5), color=S.INK_SOFT,
        line_spacing=1.32, margin=Inches(0),
    )
    cur_y = cur_y + body_h + Inches(0.05)

    # evidence tinted block
    ev = c.get("evidence")
    ev_label = c.get("evidence_label", "Signal in the data")
    if ev:
        ev_h = Inches(0.85)
        add_rect(slide, x + pad, cur_y, inner_w, ev_h,
                 fill=S.PURPLE_TINT, line=S.PURPLE_EDGE)
        add_paragraphs(
            slide, x + pad + Inches(0.10), cur_y + Inches(0.08),
            inner_w - Inches(0.2), ev_h - Inches(0.16),
            [
                {"runs": [
                    {"text": ev_label.upper() + "  ",
                     "size": Pt(8), "bold": True, "color": S.PURPLE,
                     "char_spacing": S.CSP_PILL},
                ], "space_after": Pt(2)},
                {"runs": [
                    {"text": ev, "size": Pt(9.5),
                     "italic": True, "color": S.PURPLE_DEEP},
                ], "line_spacing": 1.28},
            ],
            margin=Inches(0),
        )
        cur_y = cur_y + ev_h + Inches(0.10)

    # implication / value line at bottom
    impl = c.get("implication", c.get("value", ""))
    impl_label = c.get("implication_label", "Implication")
    if impl:
        rem_h = y + h - cur_y - Inches(0.15)
        if rem_h < Inches(0.5):
            rem_h = Inches(0.5)
        add_rect(slide, x + pad, cur_y, Inches(0.5), Inches(0.02),
                 fill=S.PURPLE, line=S.PURPLE)
        add_paragraphs(
            slide, x + pad, cur_y + Inches(0.10), inner_w, rem_h,
            [{"runs": [
                {"text": impl_label + ":  ", "size": Pt(9),
                 "bold": True, "italic": True, "color": S.PURPLE},
                {"text": impl, "size": Pt(9.5),
                 "italic": True, "color": S.INK_SOFT},
            ], "line_spacing": 1.28}],
            margin=Inches(0),
        )


# ---------- Action Table ----------
def build_action_table(slide, sd, page, total):
    add_header(slide, sd.get("eyebrow", ""), sd.get("title", ""),
               sd.get("subtitle"))
    add_footer(slide, page, total)

    actions = sd.get("actions", [])
    if not actions:
        return

    # Custom 4-col layout: # | Action | What | Value
    total_w = Inches(12.333)
    x0 = Inches(0.5)
    y0 = Inches(2.20)
    col_ws = [Inches(0.55), Inches(3.2), Inches(4.7), Inches(3.883)]
    headers = ["#", "ACTION", "WHAT TO DO", "EXPECTED VALUE"]

    header_h = Inches(0.42)
    add_rect(slide, x0, y0, total_w, header_h, fill=S.INK, line=S.INK)
    cur_x = x0
    for i, h in enumerate(headers):
        align = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
        add_textbox(
            slide, cur_x + (Inches(0) if i == 0 else Inches(0.15)),
            y0, col_ws[i] - (Inches(0) if i == 0 else Inches(0.15)),
            header_h,
            h,
            font_size=Pt(10.5), bold=True, color=S.WHITE,
            char_spacing=S.CSP_PILL, align=align,
            anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0.05),
        )
        cur_x = cur_x + col_ws[i]

    # body rows
    n_rows = len(actions)
    body_h = Inches(6.95) - (y0 + header_h)
    row_h = body_h / max(n_rows, 1)
    if row_h > Inches(0.95):
        row_h = Inches(0.95)

    cur_y = y0 + header_h
    for r_idx, act in enumerate(actions):
        row_fill = S.BG_SOFT if r_idx % 2 == 1 else S.WHITE
        add_rect(slide, x0, cur_y, total_w, row_h,
                 fill=row_fill, line=S.BORDER)
        # # column - purple
        add_textbox(
            slide, x0, cur_y, col_ws[0], row_h,
            f"{r_idx+1:02d}",
            font_size=Pt(16), bold=True, color=S.PURPLE,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0),
        )
        # Action header
        add_textbox(
            slide, x0 + col_ws[0] + Inches(0.15), cur_y,
            col_ws[1] - Inches(0.20), row_h,
            act.get("header", ""),
            font_size=Pt(10), bold=True, color=S.INK,
            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.22,
            margin=Inches(0.02),
        )
        # What
        add_textbox(
            slide, x0 + col_ws[0] + col_ws[1] + Inches(0.15), cur_y,
            col_ws[2] - Inches(0.20), row_h,
            act.get("what", ""),
            font_size=Pt(9), color=S.INK_SOFT,
            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.25,
            margin=Inches(0.02),
        )
        # Value
        add_textbox(
            slide, x0 + col_ws[0] + col_ws[1] + col_ws[2] + Inches(0.15), cur_y,
            col_ws[3] - Inches(0.20), row_h,
            act.get("value", ""),
            font_size=Pt(9), italic=True, color=S.PURPLE_DEEP,
            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.25,
            margin=Inches(0.02),
        )
        cur_y = cur_y + row_h


# ---------- Summary (pillars) ----------
def build_summary(slide, sd, page, total):
    add_header(slide, sd.get("eyebrow", ""), sd.get("title", ""),
               sd.get("subtitle"))
    add_footer(slide, page, total)

    pillars = sd.get("pillars", sd.get("takeaways", []))[:3]
    n = len(pillars)
    if n == 0:
        return

    total_w = Inches(12.333)
    gap = Inches(0.20)
    card_w = (total_w - gap * (n - 1)) / n
    card_h = Inches(3.75)
    top = Inches(2.30)

    for i, p in enumerate(pillars):
        x = Inches(0.5) + (card_w + gap) * i
        add_card(slide, x, top, card_w, card_h, accent_top=True)
        # number
        num = f"{i+1:02d}"
        add_textbox(
            slide, x, top + Inches(0.30), card_w, Inches(0.85),
            num,
            font_size=S.SZ_AGENDA_NUM, bold=True, color=S.PURPLE_DEEP,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0.05),
        )
        # rule
        rule_w = Inches(0.5)
        add_rect(slide, x + (card_w - rule_w) / 2, top + Inches(1.35),
                 rule_w, Inches(0.025), fill=S.PURPLE, line=S.PURPLE)
        # header
        add_textbox(
            slide, x + Inches(0.20), top + Inches(1.50),
            card_w - Inches(0.40), Inches(0.55),
            p.get("header", p.get("title", "")),
            font_size=Pt(13), bold=True, color=S.INK,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, margin=Inches(0.02),
        )
        # body
        add_textbox(
            slide, x + Inches(0.22), top + Inches(2.15),
            card_w - Inches(0.44), Inches(1.5),
            p.get("body", ""),
            font_size=Pt(10), italic=True, color=S.INK_SOFT,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP,
            line_spacing=1.32, margin=Inches(0.02),
        )

    # bottom takeaway strip
    tail = sd.get("closing_line")
    if tail:
        strip_y = Inches(6.40)
        add_rect(slide, Inches(0.5), strip_y, Inches(12.333), Inches(0.55),
                 fill=S.INK, line=S.INK)
        add_rect(slide, Inches(0.5), strip_y, Inches(0.12), Inches(0.55),
                 fill=S.PURPLE, line=S.PURPLE)
        add_textbox(
            slide, Inches(0.75), strip_y, Inches(11.9), Inches(0.55),
            tail,
            font_size=Pt(11.5), bold=True, italic=True, color=S.WHITE,
            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.25, margin=Inches(0.05),
        )


# ---------- References (primary / method / citations) ----------
def build_references(slide, sd, page, total):
    add_header(slide, sd.get("eyebrow", ""), sd.get("title", ""),
               sd.get("subtitle"))
    add_footer(slide, page, total)

    # 3 columns: primary | method | citations
    blocks = [sd.get("primary"), sd.get("method"), sd.get("citations")]
    blocks = [b for b in blocks if b]
    if not blocks:
        return

    total_w = Inches(12.333)
    gap = Inches(0.20)
    n = len(blocks)
    col_w = (total_w - gap * (n - 1)) / n
    col_h = Inches(4.85)
    top = Inches(2.15)

    for i, b in enumerate(blocks):
        x = Inches(0.5) + (col_w + gap) * i
        add_card(slide, x, top, col_w, col_h, accent_top=True)
        pad = Inches(0.25)
        cur_y = top + Inches(0.22)
        # header
        add_textbox(
            slide, x + pad, cur_y, col_w - pad * 2, Inches(0.45),
            b.get("header", "").upper(),
            font_size=Pt(11), bold=True, color=S.PURPLE,
            char_spacing=S.CSP_EYEBROW, margin=Inches(0),
        )
        cur_y = cur_y + Inches(0.50)
        # rule
        add_rect(slide, x + pad, cur_y, Inches(0.5), Inches(0.02),
                 fill=S.PURPLE, line=S.PURPLE)
        cur_y = cur_y + Inches(0.15)
        # items
        items = b.get("items", [])
        specs = []
        for it in items:
            specs.append({
                "runs": [
                    {"text": S.BULLET_CHAR + "  ", "size": Pt(8),
                     "color": S.PURPLE, "bold": True},
                    {"text": it, "size": Pt(9), "color": S.INK_SOFT},
                ],
                "line_spacing": 1.30,
                "space_after": Pt(5),
            })
        if specs:
            add_paragraphs(
                slide, x + pad, cur_y, col_w - pad * 2,
                top + col_h - cur_y - Inches(0.15),
                specs, margin=Inches(0),
            )


# ---------- Closing ----------
def build_closing(slide, sd, page, total):
    add_header(slide, sd.get("eyebrow", ""), sd.get("title", ""),
               sd.get("subtitle"))
    add_footer(slide, page, total)

    url_label = sd.get("url_label", "Read the full report")
    url_text = sd.get("url_text", "")

    # centered card
    card_w = Inches(10.0)
    card_h = Inches(3.6)
    x = (S.SLIDE_W - card_w) / 2
    y = Inches(2.85)
    add_card(slide, x, y, card_w, card_h, accent_top=True)

    # big "Thank You" already in header; add CTA + URL
    add_textbox(
        slide, x + Inches(0.5), y + Inches(0.70),
        card_w - Inches(1.0), Inches(0.6),
        url_label,
        font_size=Pt(18), bold=True, color=S.INK,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0),
    )
    # purple rule
    rule_w = Inches(0.8)
    add_rect(slide, x + (card_w - rule_w) / 2, y + Inches(1.45),
             rule_w, Inches(0.03), fill=S.PURPLE, line=S.PURPLE)
    # url
    if url_text:
        add_textbox(
            slide, x + Inches(0.5), y + Inches(1.70),
            card_w - Inches(1.0), Inches(0.5),
            url_text,
            font_size=Pt(12), italic=True, color=S.PURPLE_DEEP,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0),
        )
    # tagline
    add_textbox(
        slide, x + Inches(0.5), y + card_h - Inches(0.85),
        card_w - Inches(1.0), Inches(0.4),
        "Curated executive briefs  ·  aidigest.club / reads",
        font_size=Pt(10), italic=True, color=S.MUTED,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0),
    )
